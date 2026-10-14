from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "apps") not in sys.path:
    sys.path.insert(0, str(ROOT / "apps"))

from lcl_bot.processor import PackingBoxProcessor  # noqa: E402


class LclPackMergeTest(unittest.TestCase):
    def test_same_sku_across_shipments_is_not_duplicated(self):
        invoice = pd.DataFrame(
            {
                "发货单号": ["SP1", None, "SP2"],
                "SKU": [801801, 801612, 801801],
                "发货数量": [120, 240, 1000],
            }
        )
        invoice["发货单号"] = invoice["发货单号"].ffill()
        pack = pd.DataFrame(
            {
                "发货单号": ["SP1", None, "SP2"],
                "SKU": [801801, 801612, 801801],
                "总重量（kg）-箱子": [13.5, 16.8, 94.5],
                "总体积（m³）-箱子": [0.03, 0.02, 0.20],
                "箱子毛重（kg）": [13.5, 16.8, 13.5],
                "单箱数量": [160, 120, 160],
                "箱子长度（cm）": [38, 30, 38],
                "箱子宽度（cm）": [30, 25, 30],
                "箱子高度（cm）": [25, 20, 25],
            }
        )

        merged = PackingBoxProcessor._merge_invoice_with_pack_info(invoice, pack)

        self.assertEqual(len(merged), 3)
        sku_rows = merged[merged["SKU"] == 801801]
        self.assertEqual(len(sku_rows), 2)
        self.assertCountEqual(sku_rows["发货单号"].tolist(), ["SP1", "SP2"])
        self.assertCountEqual(sku_rows["发货数量"].tolist(), [120, 1000])

    def test_load_keeps_one_row_per_warehouse_for_shared_sku(self):
        invoice = pd.DataFrame(
            {
                "发货单号": ["SP260916040", None, "SP260916039"],
                "发货仓库（单据）": ["青山湖仓库", None, "良品仓"],
                "品名": ["3pc地毯割刀OMT", "other", "3pc地毯割刀OMT"],
                "SKU": [801801, 801612, 801801],
                "包装规格": ["13.50x1.00x8.50", "16.00x11.80x3.00", "13.50x1.00x8.50"],
                "单品毛重": [80.0, 140.0, 80.0],
                "发货量": [120, 240, 1000],
            }
        )
        pack = pd.DataFrame(
            {
                "发货单号": ["SP260916040", None, "SP260916039"],
                "SKU": [801801, 801612, 801801],
                "发货数量": [120, 240, 1000],
                "单箱数量": [160, 120, 160],
                "箱子毛重（kg）": [13.5, 16.8, 13.5],
                "箱子长度（cm）": [38, 30, 38],
                "箱子宽度（cm）": [30, 25, 30],
                "箱子高度（cm）": [25, 20, 25],
                "CBM（m³）-箱子": [0.03, 0.02, 0.20],
                "总重量（kg）-箱子": [13.5, 16.8, 94.5],
            }
        )

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "shipment.xlsx"
            with pd.ExcelWriter(path) as writer:
                invoice.to_excel(writer, sheet_name="发货单详情", index=False)
                pack.to_excel(writer, sheet_name="装箱信息", index=False)

            processor = PackingBoxProcessor(str(path), str(Path(tmp) / "out.xlsx"))
            loaded = processor._load_and_preprocess_data()

        sku_rows = loaded[loaded["SKU"].astype(str) == "801801"]
        self.assertEqual(len(sku_rows), 2)
        self.assertCountEqual(sku_rows["发货仓库（单据）"].tolist(), ["青山湖仓库", "良品仓"])
        self.assertCountEqual(sku_rows["发货数量"].tolist(), [120, 1000])

    def test_collapse_keeps_same_sku_from_two_groups(self):
        expanded = pd.DataFrame(
            {
                "SKU": ["801801", "801801", "801612"],
                "小组名称": ["良品仓-第1组", "青山湖仓库-第2组", "良品仓-第3组"],
                "良品仓-第1组单份数量1": [125, "", ""],
                "青山湖仓库-第2组单份数量1": ["", 24, ""],
                "良品仓-第3组单份数量1": ["", "", 48],
            }
        )
        amazon = pd.DataFrame({"SKU": ["801801", "801612"]})
        out = PackingBoxProcessor._collapse_sku_box_columns(expanded, amazon)
        self.assertEqual(list(out["SKU"]), ["801801", "801612"])
        self.assertEqual(out.loc[0, "良品仓-第1组单份数量1"], 125)
        self.assertEqual(out.loc[0, "青山湖仓库-第2组单份数量1"], 24)
        self.assertEqual(out.loc[1, "良品仓-第3组单份数量1"], 48)

    def test_floor_non_integer_share_quantity_and_total(self):
        processor = PackingBoxProcessor("in.xlsx", "out.xlsx")
        df = pd.DataFrame(
            {
                "SKU": ["80375221", "801801", "tiny"],
                "发货数量": [74, 1000, 4],
                "单箱数量": [60, 125, 10],
                "箱数": [74 / 60, 8.0, 0.4],
                "是否拼箱": ["拼箱", "拼箱", "拼箱"],
                "_force_packing": [False, True, False],
            }
        )

        out = processor._floor_non_integer_share_quantities(df)
        hit = out.loc[out["SKU"] == "80375221"].iloc[0]
        self.assertEqual(hit["发货数量"], 70)
        self.assertEqual(processor._determine_units_per_box_for_packing(hit), 14)

        forced = out.loc[out["SKU"] == "801801"].iloc[0]
        self.assertEqual(forced["发货数量"], 1000)
        self.assertEqual(processor._determine_units_per_box_for_packing(forced), 125)

        tiny = out.loc[out["SKU"] == "tiny"].iloc[0]
        self.assertEqual(tiny["发货数量"], 4)

    def test_packing_result_columns_drop_internal_fields(self):
        processor = PackingBoxProcessor("in.xlsx", "out.xlsx")
        cols = processor._get_output_columns(include_warehouse=True)
        self.assertEqual(
            cols,
            [
                "SKU",
                "发货仓库（单据）",
                "发货数量",
                "是否拼箱",
                "单份数量",
                "小组名称",
                "实际箱数",
                "单箱重量",
                "单箱理论长",
                "单箱理论宽",
                "单箱理论高",
            ],
        )
        dropped = {
            "品名",
            "箱数",
            "单品毛重（g）",
            "箱子毛重（kg）",
            "单品长（cm）",
            "单品宽（cm）",
            "单品高（cm）",
            "备注",
            "单份总重量",
            "重量差",
            "单品体积",
            "单份体积",
            "单箱体积",
        }
        self.assertTrue(dropped.isdisjoint(cols))

    def test_resolve_amazon_packaging_sheet_english_and_chinese(self):
        self.assertEqual(
            PackingBoxProcessor.resolve_amazon_packaging_sheet(
                ["Instructions", "Box packing information", "Metadata"]
            ),
            "Box packing information",
        )
        self.assertEqual(
            PackingBoxProcessor.resolve_amazon_packaging_sheet(["包装箱包装信息"]),
            "包装箱包装信息",
        )
        with self.assertRaises(ValueError) as ctx:
            PackingBoxProcessor.resolve_amazon_packaging_sheet(["Instructions", "Metadata"])
        self.assertIn("Box packing information", str(ctx.exception))

    def test_insert_amazon_packaging_english_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            amazon_path = Path(tmp) / "en.xlsx"
            _write_sta_packaging_stub(
                amazon_path,
                sheet_name="Box packing information",
                skus=[("804016", 700), ("804020", 400)],
                labels=("Name of box", "Box weight (lb):", "Box width (inch):", "Box length (inch):", "Box height (inch):"),
            )
            merge = pd.DataFrame(
                {
                    "SKU": ["804020", "804016"],
                    "发货数量": [400, 700],
                    "单份数量": [80, 100],
                    "小组名称": ["杭州虚拟仓-第1组", "杭州虚拟仓-第3组"],
                    "实际箱数": [5, 7],
                    "单箱重量": [17975.0, 12640.0],
                    "单箱理论长": [32.0, 29.0],
                    "单箱理论宽": [25.0, 38.0],
                    "单箱理论高": [54.94, 31.0],
                }
            )
            processor = PackingBoxProcessor("in.xlsx", "out.xlsx")
            processor.insert_data_to_amazon_packaging(str(amazon_path), merge)

            wb = load_workbook(amazon_path, data_only=True)
            ws = wb["Box packing information"]
            self.assertEqual(ws.cell(5, 13).value, "Box 1 quantity")
            self.assertEqual(ws.cell(9, 13).value, "P1 - B1")
            self.assertIsNotNone(ws.cell(10, 13).value)
            self.assertEqual(ws.cell(6, 11).value, 700)
            # 第1组 5 箱在前，第3组 7 箱在后；804016 应写在 M+5 起
            self.assertEqual(ws.cell(6, 18).value, 100)
            self.assertEqual(ws.cell(7, 13).value, 80)
            wb.close()


def _write_sta_packaging_stub(path: Path, *, sheet_name: str, skus: list[tuple[str, int]], labels: tuple[str, ...]) -> None:
    wb = Workbook()
    wb.active.title = "Instructions"
    ws = wb.create_sheet(sheet_name)
    ws["A3"] = "Total SKUs"
    ws["I3"] = "Total box count:"
    ws["M3"] = 0
    headers = [
        "SKU",
        "Product title ",
        "Id",
        "ASIN",
        "FNSKU",
        "Condition",
        "Prep type",
        "Who preps units?",
        "Who labels units?",
        "Expected quantity",
        "Boxed quantity",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(5, col, header)
    for offset, (sku, qty) in enumerate(skus):
        ws.cell(6 + offset, 1, sku)
        ws.cell(6 + offset, 10, qty)
    spec_row = 6 + len(skus) + 1
    for i, label in enumerate(labels):
        ws.cell(spec_row + i, 1, label)
    wb.create_sheet("Metadata")
    wb.save(path)


if __name__ == "__main__":
    unittest.main()
